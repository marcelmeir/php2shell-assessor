import csv
from pathlib import Path

from openpyxl import load_workbook

from php2shell.exporters import export_results
from php2shell.models import HostResult, ServiceResult
from php2shell.storage import ResultStore


def test_all_exports_and_workbook_sheets(tmp_path: Path):
    store = ResultStore(tmp_path / "checkpoint.sqlite3")
    service = ServiceResult(
        ip_address="192.0.2.5",
        hostname="app.example.edu",
        reverse_dns="host.example.edu",
        port=443,
        protocol="HTTPS",
        url="https://app.example.edu:443/",
        php_detected=True,
        php_confidence="CONFIRMED",
        php_version="8.3.35",
        required_patch_level="8.3.36",
        security_assessment="CRITICAL",
        suspected_vulnerability="YES",
        recommended_action="Verify and update",
        detection_evidence="X-Powered-By contains PHP",
    )
    store.save_host(
        HostResult(
            ip_address="192.0.2.5", reachable=True, web_service_count=1,
            status="WEB_SERVER_FOUND"
        ),
        [service],
    )
    files = export_results(store, tmp_path / "results")
    store.close()
    assert all(path.exists() for path in files.values())
    with files["php_critical"].open(encoding="utf-8-sig") as handle:
        assert next(csv.DictReader(handle))["php_version"] == "8.3.35"
    workbook = load_workbook(files["excel"])
    assert workbook.sheetnames == [
        "Dashboard",
        "Kritische Systeme",
        "PHP erkannt",
        "Manuelle Prüfung",
        "Alle Webserver",
        "Alle Hosts",
        "Scanfehler",
    ]
    assert workbook["Alle Webserver"].freeze_panes == "A2"
    assert workbook["Alle Webserver"].auto_filter.ref
