from __future__ import annotations

import csv
from dataclasses import fields
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import HostResult, ServiceResult
from .storage import ResultStore

RED = "F4CCCC"
ORANGE = "FCE5CD"
YELLOW = "FFF2CC"
GREEN = "D9EAD3"
GRAY = "E7E6E6"
NAVY = "1F4E78"
WHITE = "FFFFFF"


def export_results(store: ResultStore, output_dir: str | Path) -> dict[str, Path]:
    output = Path(output_dir)
    output.mkdir(parents=True, exist_ok=True)
    hosts = store.hosts()
    services = store.services()
    php = [row for row in services if row["php_confidence"] != "UNKNOWN"]
    critical = [row for row in services if row["security_assessment"] == "CRITICAL"]
    manual = [
        row
        for row in services
        if row["security_assessment"] in {"PHP_VERSION_UNKNOWN", "PHP_SUSPECTED"}
        or (
            row["php_confidence"] != "UNKNOWN"
            and row["verified_patch_status"] == "NOT_VERIFIED"
        )
    ]
    errors = [row for row in hosts if row["status"] == "SCAN_ERROR" or row["error"]]
    errors.extend(row for row in services if row["errors"])

    files = {
        "all_hosts": output / "all_hosts.csv",
        "webservers": output / "webservers.csv",
        "php_detected": output / "php_detected.csv",
        "php_critical": output / "php_critical.csv",
        "manual_review": output / "manual_review.csv",
        "excel": output / "scan_results.xlsx",
    }
    host_fields = [item.name for item in fields(HostResult)]
    service_fields = [item.name for item in fields(ServiceResult)]
    _write_csv(files["all_hosts"], hosts, host_fields)
    _write_csv(files["webservers"], services, service_fields)
    _write_csv(files["php_detected"], php, service_fields)
    _write_csv(files["php_critical"], critical, service_fields)
    _write_csv(files["manual_review"], manual, service_fields)
    _write_excel(files["excel"], hosts, services, php, critical, manual, errors)
    return files


def _write_csv(
    path: Path, rows: list[dict[str, Any]], fieldnames: list[str]
) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def _write_excel(
    path: Path,
    hosts: list[dict[str, Any]],
    services: list[dict[str, Any]],
    php: list[dict[str, Any]],
    critical: list[dict[str, Any]],
    manual: list[dict[str, Any]],
    errors: list[dict[str, Any]],
) -> None:
    workbook = Workbook()
    dashboard = workbook.active
    dashboard.title = "Dashboard"
    dashboard.sheet_view.showGridLines = False
    dashboard["A1"] = "PHP2Shell Network Discovery & Vulnerability Assessment"
    dashboard["A1"].font = Font(size=16, bold=True, color=WHITE)
    dashboard["A1"].fill = PatternFill("solid", fgColor=NAVY)
    dashboard.merge_cells("A1:D1")

    metrics = [
        ("Untersuchte IP-Adressen", len(hosts), GRAY),
        ("Erreichbare Webdienste", len(services), GRAY),
        ("PHP erkannt/vermutet", len(php), YELLOW),
        ("Potenziell verwundbar", len(critical), RED),
        (
            "Aktueller Patchstand erkannt",
            sum(row["security_assessment"] == "PATCH_LEVEL_DETECTED" for row in services),
            GREEN,
        ),
        ("Manuelle Prüfung", len(manual), ORANGE),
        (
            "Verifiziert gepatcht (Inventar)",
            sum(row["verified_patch_status"].upper() == "PATCHED" for row in services),
            GREEN,
        ),
    ]
    for index, (label, value, color) in enumerate(metrics, start=3):
        dashboard.cell(index, 1, label)
        dashboard.cell(index, 2, value)
        dashboard.cell(index, 1).fill = PatternFill("solid", fgColor=color)
        dashboard.cell(index, 2).fill = PatternFill("solid", fgColor=color)
        dashboard.cell(index, 2).font = Font(bold=True)
    dashboard["A12"] = (
        "Hinweis: HTTP-Versionsangaben sind nicht verifiziert; Distributionen können "
        "Sicherheitsfixes zurückportieren. Verifiziert wird nur durch vertrauenswürdige Inventardaten."
    )
    dashboard.merge_cells("A12:D14")
    dashboard["A12"].alignment = Alignment(wrap_text=True, vertical="top")
    dashboard.column_dimensions["A"].width = 42
    dashboard.column_dimensions["B"].width = 22

    sheets = [
        ("Kritische Systeme", critical),
        ("PHP erkannt", php),
        ("Manuelle Prüfung", manual),
        ("Alle Webserver", services),
        ("Alle Hosts", hosts),
        ("Scanfehler", errors),
    ]
    for title, rows in sheets:
        sheet = workbook.create_sheet(title)
        _populate_sheet(sheet, rows)
    workbook.save(path)


def _populate_sheet(sheet: Any, rows: list[dict[str, Any]]) -> None:
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    if not rows:
        sheet["A1"] = "Keine Einträge"
        sheet["A1"].font = Font(bold=True, color=WHITE)
        sheet["A1"].fill = PatternFill("solid", fgColor=NAVY)
        return
    headers = list(rows[0])
    sheet.append(headers)
    for row in rows:
        sheet.append([row.get(header, "") for header in headers])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(wrap_text=True)
    sheet.auto_filter.ref = sheet.dimensions
    status_index = headers.index("security_assessment") + 1 if "security_assessment" in headers else 0
    confidence_index = headers.index("php_confidence") + 1 if "php_confidence" in headers else 0
    for row_index in range(2, sheet.max_row + 1):
        status = sheet.cell(row_index, status_index).value if status_index else ""
        confidence = sheet.cell(row_index, confidence_index).value if confidence_index else ""
        color = _row_color(str(status), str(confidence))
        if color:
            for cell in sheet[row_index]:
                cell.fill = PatternFill("solid", fgColor=color)
        for cell in sheet[row_index]:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for index, header in enumerate(headers, start=1):
        values = [str(header)] + [str(row.get(header, "")) for row in rows[:200]]
        width = min(max(len(value) for value in values) + 2, 55)
        sheet.column_dimensions[get_column_letter(index)].width = max(width, 10)


def _row_color(status: str, confidence: str) -> str:
    if status == "CRITICAL":
        return RED
    if status in {"PHP_VERSION_UNKNOWN", "PHP_SUSPECTED"}:
        return ORANGE
    if confidence in {"LOW", "MEDIUM", "HIGH"}:
        return YELLOW
    if status == "PATCH_LEVEL_DETECTED":
        return GREEN
    return GRAY if status == "WEB_SERVER_FOUND" else ""
