from __future__ import annotations

import csv
import ipaddress
import re
import socket
from collections.abc import Callable, Iterable
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

from .targets import TargetSet

DOMAIN_HEADERS = {
    "domain",
    "domainname",
    "dns",
    "dnsname",
    "fqdn",
    "host",
    "hostname",
    "ownername",
}
IP_HEADERS = {"address", "addresses", "ip", "ipaddress", "ipaddresses", "ipv4"}
STATUS_HEADERS = {"status", "dnsstatus", "state"}
REPORT_FIELDS = [
    "dns_name",
    "source_status",
    "supplied_ips",
    "resolved_ips",
    "in_scope_ips",
    "outside_scope_ips",
    "disposition",
    "resolution_error",
]


@dataclass(slots=True)
class DomainEntry:
    name: str
    supplied_ips: set[str]
    status: str = ""


@dataclass(slots=True)
class DomainPreparation:
    hostnames: dict[str, list[str]]
    total_names: int
    in_scope_names: int
    outside_scope_names: int
    unresolved_names: int
    skipped_wildcards: int
    skipped_invalid: int
    mapping_count: int
    report_path: Path
    mapping_path: Path


def prepare_domain_list(
    source: str | Path,
    targets: TargetSet,
    output_dir: str | Path,
    *,
    resolve_dns: bool = True,
    resolver_workers: int = 20,
    resolver: Callable[[str], tuple[list[str], str]] | None = None,
) -> DomainPreparation:
    """Load a generic domain list, resolve it and retain only in-scope IPv4 mappings."""
    source_path = Path(source).expanduser().resolve()
    entries, invalid_count = load_domain_entries(source_path)
    output = Path(output_dir)
    output.mkdir(parents=True, exist_ok=True)
    resolver = resolver or resolve_ipv4

    concrete: dict[str, DomainEntry] = {}
    wildcard_count = 0
    for entry in entries:
        if "*" in entry.name:
            wildcard_count += 1
            continue
        existing = concrete.get(entry.name)
        if existing:
            existing.supplied_ips.update(entry.supplied_ips)
            if not existing.status:
                existing.status = entry.status
        else:
            concrete[entry.name] = entry

    resolutions: dict[str, tuple[list[str], str]] = {}
    if resolve_dns:
        with ThreadPoolExecutor(max_workers=max(1, min(resolver_workers, 50))) as executor:
            futures = {executor.submit(resolver, name): name for name in concrete}
            for future in as_completed(futures):
                name = futures[future]
                try:
                    resolutions[name] = future.result()
                except Exception as exc:  # a custom resolver must not abort the scan setup
                    resolutions[name] = ([], str(exc))

    mappings: dict[str, set[str]] = {}
    rows: list[dict[str, str]] = []
    dispositions = {"IN_SCOPE": 0, "OUTSIDE_SCOPE": 0, "UNRESOLVED": 0}
    for name, entry in sorted(concrete.items()):
        resolved, error = resolutions.get(
            name,
            ([], "DNS resolution disabled" if not resolve_dns else "resolution not completed"),
        )
        candidates = entry.supplied_ips | set(resolved)
        in_scope = sorted(address for address in candidates if targets.contains(address))
        outside = sorted(address for address in candidates if not targets.contains(address))
        for address in in_scope:
            mappings.setdefault(address, set()).add(name)
        disposition = "IN_SCOPE" if in_scope else ("OUTSIDE_SCOPE" if outside else "UNRESOLVED")
        dispositions[disposition] += 1
        rows.append(
            {
                "dns_name": name,
                "source_status": entry.status,
                "supplied_ips": ";".join(sorted(entry.supplied_ips)),
                "resolved_ips": ";".join(resolved),
                "in_scope_ips": ";".join(in_scope),
                "outside_scope_ips": ";".join(outside),
                "disposition": disposition,
                "resolution_error": error,
            }
        )

    for entry in sorted((item for item in entries if "*" in item.name), key=lambda x: x.name):
        rows.append(
            {
                "dns_name": entry.name,
                "source_status": entry.status,
                "supplied_ips": ";".join(sorted(entry.supplied_ips)),
                "resolved_ips": "",
                "in_scope_ips": "",
                "outside_scope_ips": "",
                "disposition": "WILDCARD_SKIPPED",
                "resolution_error": "Wildcard names are not expanded",
            }
        )

    report_path = output / "domain_scope_report.csv"
    with report_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=REPORT_FIELDS)
        writer.writeheader()
        writer.writerows(rows)
    mapping_path = output / "hostnames_in_scope.txt"
    mapping_path.write_text(
        "# Generated from domain list; only mappings inside the configured target scope\n"
        + "".join(
            f"{address} {name}\n"
            for address, names in sorted(mappings.items())
            for name in sorted(names)
        ),
        encoding="utf-8",
    )
    normalized = {address: sorted(names) for address, names in mappings.items()}
    return DomainPreparation(
        hostnames=normalized,
        total_names=len(concrete),
        in_scope_names=dispositions["IN_SCOPE"],
        outside_scope_names=dispositions["OUTSIDE_SCOPE"],
        unresolved_names=dispositions["UNRESOLVED"],
        skipped_wildcards=wildcard_count,
        skipped_invalid=invalid_count,
        mapping_count=sum(len(names) for names in normalized.values()),
        report_path=report_path,
        mapping_path=mapping_path,
    )


def load_domain_entries(path: Path) -> tuple[list[DomainEntry], int]:
    suffix = path.suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        raw_rows = _rows_from_xlsx(path)
    elif suffix in {".csv", ".tsv"}:
        raw_rows = _rows_from_delimited(path)
    elif suffix in {".txt", ".list"}:
        raw_rows = _rows_from_text(path)
    else:
        raise ValueError("Domain list must be TXT, LIST, CSV, TSV, XLSX or XLSM")
    entries: list[DomainEntry] = []
    invalid = 0
    for name, addresses, status in raw_rows:
        normalized = normalize_hostname(name)
        if not normalized:
            invalid += 1
            continue
        entries.append(DomainEntry(normalized, _parse_ipv4_values(addresses), status.strip()))
    if not entries:
        raise ValueError("No valid DNS names were found in the domain list")
    return entries, invalid


def resolve_ipv4(hostname: str) -> tuple[list[str], str]:
    try:
        records = socket.getaddrinfo(hostname, None, socket.AF_INET, socket.SOCK_STREAM)
        return sorted({record[4][0] for record in records}), ""
    except OSError as exc:
        return [], str(exc)


def normalize_hostname(value: object) -> str:
    text = str(value or "").strip().rstrip(".").lower()
    if not text or len(text) > 253 or " " in text:
        return ""
    wildcard = text.startswith("*.")
    core = text[2:] if wildcard else text
    try:
        ascii_core = core.encode("idna").decode("ascii")
    except UnicodeError:
        return ""
    labels = ascii_core.split(".")
    if len(labels) < 2 or any(
        not label
        or len(label) > 63
        or label.startswith("-")
        or label.endswith("-")
        or not re.fullmatch(r"[a-z0-9-]+", label)
        for label in labels
    ):
        return ""
    return f"*.{ascii_core}" if wildcard else ascii_core


def _normalize_header(value: object) -> str:
    return re.sub(r"[^a-z0-9]", "", str(value or "").strip().lower())


def _find_column(headers: list[object], candidates: set[str]) -> int | None:
    for index, header in enumerate(headers):
        if _normalize_header(header) in candidates:
            return index
    return None


def _rows_from_xlsx(path: Path) -> list[tuple[object, object, str]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    rows: list[tuple[object, object, str]] = []
    for sheet in workbook.worksheets:
        iterator = sheet.iter_rows(values_only=True)
        header: list[object] | None = None
        domain_col: int | None = None
        ip_col: int | None = None
        status_col: int | None = None
        for row_number, values in enumerate(iterator, start=1):
            values_list = list(values)
            if header is None:
                domain_col = _find_column(values_list, DOMAIN_HEADERS)
                if domain_col is None:
                    if row_number >= 50:
                        break
                    continue
                header = values_list
                ip_col = _find_column(header, IP_HEADERS)
                status_col = _find_column(header, STATUS_HEADERS)
                continue
            if domain_col is not None and domain_col < len(values_list) and values_list[domain_col]:
                ip_value = values_list[ip_col] if ip_col is not None and ip_col < len(values_list) else ""
                status = (
                    str(values_list[status_col] or "")
                    if status_col is not None and status_col < len(values_list)
                    else ""
                )
                rows.append((values_list[domain_col], ip_value, status))
    if not rows:
        raise ValueError("No worksheet with a recognized domain/hostname column was found")
    return rows


def _rows_from_delimited(path: Path) -> list[tuple[object, object, str]]:
    text = path.read_text(encoding="utf-8-sig")
    if path.suffix.lower() == ".tsv":
        dialect = csv.excel_tab
    else:
        try:
            dialect = csv.Sniffer().sniff(text[:8192], ",;\t")
        except csv.Error:
            dialect = csv.excel
    reader = csv.reader(text.splitlines(), dialect)
    headers = next(reader, [])
    domain_col = _find_column(headers, DOMAIN_HEADERS)
    if domain_col is None:
        return _rows_from_text(path)
    ip_col = _find_column(headers, IP_HEADERS)
    status_col = _find_column(headers, STATUS_HEADERS)
    rows = []
    for values in reader:
        if domain_col >= len(values) or not values[domain_col].strip():
            continue
        ip_value = values[ip_col] if ip_col is not None and ip_col < len(values) else ""
        status = values[status_col] if status_col is not None and status_col < len(values) else ""
        rows.append((values[domain_col], ip_value, status))
    return rows


def _rows_from_text(path: Path) -> list[tuple[object, object, str]]:
    rows = []
    for raw in path.read_text(encoding="utf-8-sig").splitlines():
        line = raw.split("#", 1)[0].strip()
        if not line:
            continue
        parts = [part for part in re.split(r"[,;\s]+", line) if part]
        addresses = [part for part in parts if _is_ipv4(part)]
        names = [part for part in parts if not _is_ipv4(part)]
        if names:
            rows.append((names[0], ";".join(addresses), ""))
    return rows


def _parse_ipv4_values(value: object) -> set[str]:
    addresses: set[str] = set()
    for part in re.split(r"[,;\s]+", str(value or "")):
        if _is_ipv4(part):
            addresses.add(str(ipaddress.IPv4Address(part)))
    return addresses


def _is_ipv4(value: str) -> bool:
    try:
        return ipaddress.ip_address(value).version == 4
    except ValueError:
        return False


def merge_hostname_maps(*maps: dict[str, list[str]]) -> dict[str, list[str]]:
    merged: dict[str, set[str]] = {}
    for mapping in maps:
        for address, names in mapping.items():
            merged.setdefault(address, set()).update(names)
    return {address: sorted(names) for address, names in merged.items()}
